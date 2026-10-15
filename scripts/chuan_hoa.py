"""Chuẩn hóa file QP-KTAT.xlsx (sổ quỹ phòng KTAT) thành một bảng giao dịch thống nhất.

Chạy:  python -I -X utf8 scripts/chuan_hoa.py
Ra:    data/giao_dich.json        (dữ liệu cho webapp)
       QP-KTAT_chuan.xlsx         (bản Excel đã chuẩn hóa: Giao dịch, Tổng hợp tháng, Cần kiểm tra)
"""
import datetime as dt
import json
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

GOC = Path(__file__).resolve().parent.parent
FILE_GOC = GOC / "QP-KTAT.xlsx"
FILE_JSON = GOC / "data" / "giao_dich.json"
FILE_XLSX = GOC / "QP-KTAT_chuan.xlsx"

# Từ khóa phân nhóm — xét theo thứ tự, khớp đầu tiên thắng.
NHOM_CHI = [
    ("Hiếu hỉ, thăm hỏi", ["viếng", "vòng hoa", "đám cưới", "đám hiếu", "thăm "]),
    ("Sinh nhật", ["sinh nhật", "sn "]),
    ("Lễ, Tết", ["tết", "8/3", "20-10", "20/10", "lì xì", "mừng tuổi", "tất niên", "tất nhiên", "khai xuân", "du xuân"]),
    ("Du lịch", ["du lịch"]),
    ("KHCN, sáng kiến", ["đề tài", "sáng kiến", "khcn", "nckh", "mc tụ", "máy cắt tụ", "tụ bù", "3d", "triển lãm"]),
    ("Ứng, vay, trả", ["ứng", "vay", "gửi a"]),
    ("Liên hoan, giao lưu", ["liên hoan", "giao lưu", "chia tay", "chè", "golden", "csttd", "cstđ"]),
    ("Công tác, làm việc", ["công tác", "làm việc", "lên ", "đi ", "xe ", "hồ sơ", "hđ", "hd ", "ban ", "tv qn", "npc", "evn", "tcns", "tchc"]),
    ("Mua sắm, quà tặng", ["mua", "quà", "tặng", "áo", "ấm chén", "camera", "rượu", "hoa quả", "hoa qua", "tranh"]),
]
NHOM_THU = [
    ("Thưởng", ["thưởng", "khen"]),
    ("Chúc tết, biếu", ["tết"]),
    ("Hoàn ứng, hoàn tiền", ["hoàn"]),
    ("Ứng, vay", ["ứng tiền", "vay", " ứng"]),
    ("Thanh toán hóa đơn", ["thanh toán", "hoá đơn", "hóa đơn", "hoa don", "hđ", "hd "]),
    ("Góp quỹ, bàn giao", ["góp quỹ", "bàn giao"]),
]

# Gộp các cách viết khác nhau của cùng một đơn vị.
DON_VI = {"npsc": "NPSC", "atuyen": "A Tuyển", "aduong": "A Dương", "dl hoa binh": "ĐL Hòa Bình"}


def phan_nhom(noi_dung: str, bang) -> str:
    s = " " + noi_dung.lower() + " "
    for ten, tu in bang:
        if any(t in s for t in tu):
            return ten
    return "Khác"


def chu(v) -> str:
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v)).strip()


def ngay(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    return None


def tien(v):
    if v is None or v == "":
        return None
    return int(round(float(v)))


def hinh_thuc(v: str) -> str:
    v = v.lower()
    if "chuyển" in v or v == "ck":
        return "CK"
    if "mặt" in v or v == "tm":
        return "TM"
    return ""


def da_hoan(*cac_cot) -> bool:
    return any("đã hoàn" in c.lower() for c in cac_cot)


def doc():
    wb = openpyxl.load_workbook(FILE_GOC, data_only=False)
    ds, loi = [], []

    ws = wb["Thu"]
    for i, r in enumerate(ws.iter_rows(min_row=2, max_col=10, values_only=True), start=2):
        stt, don_vi, ng, nd, so, ghi_chu, luu_y = r[:7]
        if all(x is None for x in r[:7]):
            continue
        nd, ghi_chu, luu_y = chu(nd), chu(ghi_chu), chu(luu_y)
        dv = chu(don_vi)
        dv = DON_VI.get(dv.lower(), dv)
        g = {
            "id": f"thu-{i:03d}",
            "loai": "thu",
            "ngay": ngay(ng),
            "soTien": tien(so),
            "noiDung": nd,
            "doiTuong": dv,
            "ghiChu": ghi_chu,
            "daHoan": da_hoan(luu_y, nd),
            "hinhThuc": "",
            "nhom": phan_nhom(nd, NHOM_THU),
            "nguon": f"Thu!{i}",
        }
        ds.append(g)
        if g["soTien"] is None:
            loi.append((g["nguon"], nd, "Không có số tiền"))
        if g["ngay"] is None:
            loi.append((g["nguon"], nd, "Không có ngày"))

    ws = wb["Chi"]
    for i, r in enumerate(ws.iter_rows(min_row=2, max_col=9, values_only=True), start=2):
        stt, ng, so, nd, ghi_chu, luu_y, tmck = r[:7]
        if all(x is None for x in r[:7]):
            continue
        nd, ghi_chu, luu_y, tmck = chu(nd), chu(ghi_chu), chu(luu_y), chu(tmck)
        g = {
            "id": f"chi-{i:03d}",
            "loai": "chi",
            "ngay": ngay(ng),
            "soTien": tien(so),
            "noiDung": nd,
            "doiTuong": "",
            "ghiChu": ghi_chu,
            "daHoan": da_hoan(luu_y, nd),
            "hinhThuc": hinh_thuc(tmck),
            "nhom": phan_nhom(nd, NHOM_CHI),
            "nguon": f"Chi!{i}",
        }
        ds.append(g)
        if g["soTien"] is None:
            loi.append((g["nguon"], nd, "Không có số tiền"))
        if g["ngay"] is None:
            loi.append((g["nguon"], nd, "Không có ngày"))
        if not g["hinhThuc"]:
            loi.append((g["nguon"], nd, "Chưa ghi TM/CK"))

    tong_goc = wb["Tổng Thu-Chi"]["A2"].value
    if isinstance(tong_goc, str) and "E118" in tong_goc:
        loi.append(("Tổng Thu-Chi!A2", tong_goc, "Công thức cộng cố định đến dòng 118 — thêm dòng sau 118 sẽ bị bỏ sót"))

    ds.sort(key=lambda g: (g["ngay"] or "", g["loai"], g["id"]))
    return ds, loi


def ghi_xlsx(ds, loi):
    wb = openpyxl.Workbook()
    dam = Font(bold=True, color="FFFFFF")
    nen = PatternFill("solid", fgColor="1F4E78")
    so_fmt = "#,##0"

    ws = wb.active
    ws.title = "Giao dịch"
    tieu_de = ["STT", "Ngày", "Loại", "Nhóm", "Nội dung", "Đơn vị / Người", "Thu", "Chi", "Tồn quỹ",
               "Hình thức", "Đã hoàn", "Ghi chú", "Dòng gốc"]
    ws.append(tieu_de)
    ton = 0
    for k, g in enumerate(ds, start=1):
        so = g["soTien"] or 0
        ton += so if g["loai"] == "thu" else -so
        ws.append([
            k,
            dt.date.fromisoformat(g["ngay"]) if g["ngay"] else None,
            "Thu" if g["loai"] == "thu" else "Chi",
            g["nhom"], g["noiDung"], g["doiTuong"],
            g["soTien"] if g["loai"] == "thu" else None,
            g["soTien"] if g["loai"] == "chi" else None,
            ton,
            {"CK": "Chuyển khoản", "TM": "Tiền mặt"}.get(g["hinhThuc"], ""),
            "x" if g["daHoan"] else "",
            g["ghiChu"], g["nguon"],
        ])
    n = len(ds) + 1
    for c in "GHI":
        for cell in ws[c][1:]:
            cell.number_format = so_fmt
    for cell in ws["B"][1:]:
        cell.number_format = "dd/mm/yyyy"
    ws.append([])
    ws.append(["", "", "", "", "TỔNG", "", f"=SUM(G2:G{n})", f"=SUM(H2:H{n})", f"=G{n+2}-H{n+2}"])
    for c in "GHI":
        ws[f"{c}{n+2}"].number_format = so_fmt
        ws[f"{c}{n+2}"].font = Font(bold=True)
    ws[f"E{n+2}"].font = Font(bold=True)
    rong = [6, 12, 6, 20, 50, 22, 14, 14, 14, 13, 8, 28, 10]
    for i, w in enumerate(rong, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:M{n}"

    # Tổng hợp theo tháng
    th = wb.create_sheet("Tổng hợp tháng")
    th.append(["Tháng", "Thu", "Chi", "Chênh lệch", "Tồn cuối tháng"])
    thang = {}
    for g in ds:
        if not g["ngay"]:
            continue
        t = g["ngay"][:7]
        a = thang.setdefault(t, [0, 0])
        a[0 if g["loai"] == "thu" else 1] += g["soTien"] or 0
    ton = 0
    for t in sorted(thang):
        thu, chi = thang[t]
        ton += thu - chi
        th.append([f"{t[5:]}/{t[:4]}", thu, chi, thu - chi, ton])
    for row in th.iter_rows(min_row=2, min_col=2, max_col=5):
        for cell in row:
            cell.number_format = so_fmt
    for c, w in zip("ABCDE", [10, 16, 16, 16, 16]):
        th.column_dimensions[c].width = w

    # Danh sách cần kiểm tra
    kt = wb.create_sheet("Cần kiểm tra")
    kt.append(["Dòng gốc", "Nội dung", "Vấn đề"])
    for l in loi:
        kt.append(list(l))
    for c, w in zip("ABC", [16, 50, 60]):
        kt.column_dimensions[c].width = w

    for s in wb.worksheets:
        for cell in s[1]:
            cell.font, cell.fill = dam, nen
            cell.alignment = Alignment(horizontal="center", vertical="center")
    wb.save(FILE_XLSX)


def main():
    ds, loi = doc()
    FILE_JSON.parent.mkdir(exist_ok=True)
    FILE_JSON.write_text(json.dumps(ds, ensure_ascii=False, indent=1), encoding="utf-8")
    ghi_xlsx(ds, loi)

    thu = sum(g["soTien"] or 0 for g in ds if g["loai"] == "thu")
    chi = sum(g["soTien"] or 0 for g in ds if g["loai"] == "chi")
    print(f"Giao dịch: {len(ds)}  (thu {sum(g['loai']=='thu' for g in ds)}, chi {sum(g['loai']=='chi' for g in ds)})")
    print(f"Tổng thu {thu:,}  Tổng chi {chi:,}  Tồn {thu-chi:,}")
    print(f"Cần kiểm tra: {len(loi)} mục")
    if "-v" in sys.argv:
        for g in ds:
            print(g["loai"], g["nhom"].ljust(22), g["noiDung"])


if __name__ == "__main__":
    main()
